import hashlib
import mimetypes
import re
from collections.abc import Callable
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation

from ..schemas import ContentUnit

MAX_CHUNK_CHARS = 1400
MEDIA_EXTENSIONS = frozenset(
    {
        ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg",
        ".mp4", ".mov", ".mkv", ".webm", ".avi", ".mp3", ".wav", ".m4a",
    }
)

_MEDIA_TECHNICAL_TOKENS = re.compile(
    r"(?ix)\b(?:"
    r"ddp\d*(?:\.\d+)?|h\.?\d{3}|\d{3,4}(?:p|i)|\d{3,4}x\d{3,4}|"
    r"\d+(?:\.\d+)?ch|\d+\.\d+|"
    r"nf|amzn|web[- ]?dl|webrip|web|dl|bluray|blu[- ]?ray|brrip|bdrip|dvdrip|hdtv|"
    r"x264|x265|h264|h265|h|264|265|hevc|avc|10bit|8bit|hdr10plus|hdr10|hdr|dv|"
    r"ddp|dd|aac|ac3|dts|truehd|atmos|multi|dual[- ]?audio|esub|remux|repack|"
    r"proper|extended|uncut|unrated|imax|rarbg|hdhub4u|4khdhub|qrips|smg|"
    r"cinemagic|ds4k|bcore|hdhub|com"
    r")\b"
)
_MEDIA_SERIES_MARKER = re.compile(r"(?ix)\b(?:s\d{1,2}e\d{1,3}|s\d{1,2})\b")
_MEDIA_RELEASE_BOUNDARY = re.compile(
    r"(?ix)\b(?:\d{3,4}(?:p|i)|web[- ]?dl|webrip|bluray|blu[- ]?ray|brrip|bdrip|dvdrip|hdtv|"
    r"x264|x265|h264|h265|hevc|avc|10bit|8bit|hdr10plus|hdr10|hdr|imax|remux)\b"
)

def extension_for(path: Path) -> str:
    return path.suffix.lower()


def mime_type_for(path: Path) -> str:
    return mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def normalize_media_filename(path: Path) -> str:
    """Keep likely title words while dropping common release metadata."""

    value = path.stem.replace("_", " ")
    value = _MEDIA_SERIES_MARKER.sub(" TV ", value)
    boundary = _MEDIA_RELEASE_BOUNDARY.search(value)
    if boundary:
        title = value[: boundary.start()]
        years = re.findall(r"\b(?:19|20)\d{2}\b", value[boundary.start() :])
        value = f"{title} {' '.join(years)}"
    value = _MEDIA_TECHNICAL_TOKENS.sub(" ", value)
    value = re.sub(r"[\[\](){}]", " ", value)
    value = value.replace(".", " ")
    value = re.sub(r"\s+", " ", value).strip(" -")
    return value or "Unknown media"


def _unit(unit_type: str, ordinal: int, text: str, locator: dict[str, str | int] | None = None) -> ContentUnit:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    return ContentUnit(
        unit_type=unit_type,
        ordinal=ordinal,
        content_hash=f"sha256:{digest}",
        text=normalized,
        locator=locator or {},
        extractor_version="2",
    )


def _text_units(text: str) -> list[ContentUnit]:
    paragraphs = [part.strip() for part in text.replace("\r\n", "\n").replace("\r", "\n").split("\n\n") if part.strip()]
    if not paragraphs and text.strip():
        paragraphs = [text.strip()]
    return [_unit("paragraph", index, paragraph) for index, paragraph in enumerate(paragraphs)]


def _split_text(text: str, max_chars: int) -> list[str]:
    """Split text without silently truncating it."""

    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return [""]
    if len(normalized) <= max_chars:
        return [normalized]

    pieces = [piece.strip() for piece in re.split(r"(?<=[.!?])\s+|\n+", normalized) if piece.strip()]
    chunks: list[str] = []
    current = ""

    for piece in pieces:
        if len(piece) > max_chars:
            if current:
                chunks.append(current)
                current = ""
            for start in range(0, len(piece), max_chars):
                chunks.append(piece[start : start + max_chars])
            continue
        candidate = f"{current} {piece}".strip()
        if current and len(candidate) > max_chars:
            chunks.append(current)
            current = piece
        else:
            current = candidate

    if current:
        chunks.append(current)
    return chunks


def _chunk_units(units: list[ContentUnit], max_chars: int) -> list[ContentUnit]:
    chunked: list[ContentUnit] = []
    for unit in units:
        parts = _split_text(unit.text, max_chars)
        for chunk_index, part in enumerate(parts):
            locator = dict(unit.locator)
            locator["chunk_index"] = chunk_index
            locator["chunk_count"] = len(parts)
            chunked.append(
                _unit(unit.unit_type, len(chunked), part, locator)
            )
    return chunked


def _read_text(path: Path) -> list[ContentUnit]:
    return _text_units(path.read_text(encoding="utf-8", errors="replace"))


def _read_pdf(path: Path) -> list[ContentUnit]:
    return [
        _unit("page", index, page.extract_text() or "", {"page": index + 1})
        for index, page in enumerate(PdfReader(str(path)).pages)
    ]


def _read_docx(path: Path) -> list[ContentUnit]:
    document = Document(str(path))
    units = [
        _unit("paragraph", index, paragraph.text, {"paragraph": index + 1})
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip()
    ]
    for table_index, table in enumerate(document.tables):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ]
        text = "\n".join(row for row in rows if row.strip())
        if text.strip():
            units.append(_unit("table", len(units), text, {"table": table_index + 1}))
    return units


def _shape_text(shape: object) -> list[str]:
    if getattr(shape, "has_table", False):
        table = shape.table
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ]
        return ["\n".join(row for row in rows if row.strip())]
    nested = getattr(shape, "shapes", None)
    if nested is not None:
        values: list[str] = []
        for child in nested:
            values.extend(_shape_text(child))
        return values
    text = getattr(shape, "text", "")
    return [text] if isinstance(text, str) and text.strip() else []


def _read_pptx(path: Path) -> list[ContentUnit]:
    units = []
    for index, slide in enumerate(Presentation(str(path)).slides):
        texts = [text for shape in slide.shapes for text in _shape_text(shape) if text.strip()]
        units.append(_unit("slide", index, "\n".join(texts), {"slide": index + 1}))
    return units


def _read_media(path: Path) -> list[ContentUnit]:
    return [_unit("media", 0, normalize_media_filename(path), {"extension": extension_for(path)})]


Reader = Callable[[Path], list[ContentUnit]]
READERS: dict[str, Reader] = {
    **{extension: _read_text for extension in {".txt", ".md", ".markdown", ".csv", ".json", ".py", ".yaml", ".yml"}},
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".pptx": _read_pptx,
    **{extension: _read_media for extension in MEDIA_EXTENSIONS},
}

SUPPORTED_EXTENSIONS = frozenset(READERS)


def extract_units(path: Path, max_chunk_chars: int = MAX_CHUNK_CHARS) -> list[ContentUnit]:
    """Select the reader from the extension registry and safely chunk output."""

    reader = READERS.get(extension_for(path))
    return _chunk_units(reader(path), max_chunk_chars) if reader else []
