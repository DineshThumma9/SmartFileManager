from pathlib import Path

from docx import Document
from pptx import Presentation

from smartfileman.readers import extract_units


def test_docx_table_content_is_extracted(tmp_path: Path) -> None:
    source = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Quarterly report")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Diagnosis"
    table.cell(0, 1).text = "Treatment"
    table.cell(1, 0).text = "Medical content"
    table.cell(1, 1).text = "Follow-up"
    document.save(source)

    units = extract_units(source)

    assert any(unit.unit_type == "table" for unit in units)
    assert "Medical content" in "\n".join(unit.text for unit in units)


def test_pptx_table_content_is_extracted(tmp_path: Path) -> None:
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, 0, 0, 5000000, 2000000).table
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Level"
    table.cell(1, 0).text = "Neo4j"
    table.cell(1, 1).text = "Advanced"
    presentation.save(source)

    units = extract_units(source)

    assert any(unit.unit_type == "slide" and "Neo4j" in unit.text for unit in units)
